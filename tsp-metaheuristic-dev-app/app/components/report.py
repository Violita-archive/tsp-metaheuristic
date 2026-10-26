"""
app/components/report.py
Ekspor hasil run, perbandingan run, dan analisis parameter.

Struktur `run` (dibuat main.py):
    label, init_method, meta_method, move, two_opt, params (dict),
    init_route, init_dist, meta_dist, final_route, final_dist, gain, pct,
    history, stop_reason, iterations, time
"""
import os
import pandas as pd

try:
    from components.export import route_table
    from components.map_plot import save_route_png, rotate_to_depot
except ImportError:
    from app.components.export import route_table
    from app.components.map_plot import save_route_png, rotate_to_depot


def params_text(params):
    return ", ".join(f"{k}={v}" for k, v in params.items()) if params else "-"


def pipeline_text(run):
    parts = [run["init_method"]]
    if run["meta_method"] != "—":
        parts.append(f"{run['meta_method']} ({run['move']})")
    if run["two_opt"]:
        parts.append("2-Opt")
    return " → ".join(parts)


def run_row(run):
    """Satu baris ringkasan run untuk tabel perbandingan."""
    return {
        "Run": run["label"],
        "Inisialisasi": run["init_method"],
        "Metaheuristik": run["meta_method"],
        "Move": run["move"] if run["meta_method"] != "—" else "-",
        "2-Opt": "Ya" if run["two_opt"] else "Tidak",
        "Parameter": params_text(run["params"]),
        "Jarak Awal": round(run["init_dist"], 4),
        "Setelah Metaheuristik": round(run["meta_dist"], 4) if run["meta_method"] != "—" else None,
        "Jarak Final": round(run["final_dist"], 4),
        "Gain (%)": round(run["pct"], 2),
        "Iterasi": run["iterations"],
        "Waktu (s)": round(run["time"], 4),
        "Kriteria Berhenti": run["stop_reason"],
    }


def _route_string(cities, route):
    r = rotate_to_depot(route)
    return " -> ".join(cities[i]["name"] for i in r + [r[0]])


def _style_book(book):
    from openpyxl.styles import Alignment, Font, PatternFill
    fill = PatternFill("solid", fgColor="FF5A1F")
    font = Font(bold=True, color="FFFFFF")
    for ws in book.worksheets:
        for cell in ws[1]:
            cell.fill, cell.font = fill, font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.freeze_panes = "A2"
        for col in ws.columns:
            width = max((len(str(c.value)) if c.value is not None else 0) for c in col[:300])
            ws.column_dimensions[col[0].column_letter].width = min(width + 2, 70)


def export_run(path, cities, run, runs=None):
    """Simpan satu run. .xlsx -> beberapa sheet, .csv -> tabel rute."""
    ext = os.path.splitext(path)[1].lower()
    table = route_table(cities, run["final_route"])
    if ext == ".csv":
        table.to_csv(path, index=False)
        return path

    summary = [
        ("Pipeline", pipeline_text(run)),
        ("Inisialisasi", run["init_method"]),
        ("Metaheuristik", run["meta_method"]),
        ("Move Operator", run["move"] if run["meta_method"] != "—" else "-"),
        ("2-Opt", "Ya" if run["two_opt"] else "Tidak"),
        ("Jumlah Kota", len(cities)),
        ("Jarak Awal", round(run["init_dist"], 4)),
        ("Setelah Metaheuristik", round(run["meta_dist"], 4)),
        ("Jarak Final", round(run["final_dist"], 4)),
        ("Gain", f"{run['gain']:.4f} ({run['pct']:.2f}%)"),
        ("Iterasi", run["iterations"]),
        ("Waktu (s)", round(run["time"], 4)),
        ("Kriteria Berhenti", run["stop_reason"]),
        ("Rute", _route_string(cities, run["final_route"])),
    ]
    with pd.ExcelWriter(path, engine="openpyxl") as w:
        pd.DataFrame(summary, columns=["Keterangan", "Nilai"]).to_excel(
            w, sheet_name="Ringkasan", index=False)
        pd.DataFrame(list(run["params"].items()) or [("-", "-")],
                     columns=["Parameter", "Nilai"]).to_excel(
            w, sheet_name="Parameter", index=False)
        table.to_excel(w, sheet_name="Rute", index=False)
        if run["history"]:
            pd.DataFrame(run["history"]).to_excel(w, sheet_name="Log Iterasi", index=False)
        if runs:
            pd.DataFrame([run_row(r) for r in runs]).to_excel(
                w, sheet_name="Riwayat Run", index=False)
        _style_book(w.book)
    return path


def export_png(path, cities, run, show_initial=True):
    coords = [(c["x"], c["y"]) for c in cities]
    names = [c["name"] for c in cities]
    save_route_png(path, coords, names, run["final_route"],
                   initial_route=run["init_route"] if show_initial else None,
                   title=f"Rute {pipeline_text(run)}",
                   subtitle=f"Jarak final: {run['final_dist']:.4f}  ·  "
                            f"awal: {run['init_dist']:.4f}  ·  gain {run['pct']:.2f}%")
    return path


def export_txt(path, cities, run, runs=None):
    line = "-" * 60
    L = [
        "TSP OPTIMIZER - LAPORAN HASIL", line,
        f"Pipeline           : {pipeline_text(run)}",
        f"Jumlah kota        : {len(cities)}",
        f"Parameter          : {params_text(run['params'])}",
        f"Jarak awal         : {run['init_dist']:.4f}",
        f"Setelah metaheur.  : {run['meta_dist']:.4f}",
        f"Jarak final        : {run['final_dist']:.4f}",
        f"Gain               : {run['gain']:.4f} ({run['pct']:.2f}%)",
        f"Iterasi            : {run['iterations']}",
        f"Waktu              : {run['time']:.4f} s",
        f"Kriteria berhenti  : {run['stop_reason']}",
        "", "Urutan kunjungan:", "  " + _route_string(cities, run["final_route"]),
    ]
    if runs and len(runs) > 1:
        L += ["", "Riwayat run:"]
        for r in runs:
            L.append(f"  {r['label']:<7} {pipeline_text(r):<45} "
                     f"{r['final_dist']:>12.4f}  ({r['pct']:.2f}%, {r['time']:.3f}s)")
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(L) + "\n")
    return path


def export_compare(path, runs):
    df = pd.DataFrame([run_row(r) for r in runs])
    if os.path.splitext(path)[1].lower() == ".csv":
        df.to_csv(path, index=False)
        return path
    with pd.ExcelWriter(path, engine="openpyxl") as w:
        df.to_excel(w, sheet_name="Perbandingan", index=False)
        conv = {}
        for r in runs:
            if r["history"]:
                conv[r["label"]] = pd.Series({h["iter"]: h["best_dist"] for h in r["history"]})
        if conv:
            pd.DataFrame(conv).rename_axis("iter").reset_index().to_excel(
                w, sheet_name="Konvergensi", index=False)
        _style_book(w.book)
    return path


def export_sweep(path, summary_rows, raw_rows, meta):
    with pd.ExcelWriter(path, engine="openpyxl") as w:
        pd.DataFrame(summary_rows).to_excel(w, sheet_name="Ringkasan", index=False)
        pd.DataFrame(raw_rows).to_excel(w, sheet_name="Data Mentah", index=False)
        pd.DataFrame(list(meta.items()), columns=["Pengaturan", "Nilai"]).to_excel(
            w, sheet_name="Pengaturan", index=False)
        _style_book(w.book)
    return path
