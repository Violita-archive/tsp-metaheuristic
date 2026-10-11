"""
app/components/export.py
Ekspor hasil TSP ke Excel (.xlsx) atau CSV (.csv).

Dipanggil main.py:
    export_results(path, cities, route, total_dist, history, method_name)
"""
import math
import os
import pandas as pd

try:
    from components.map_plot import save_route_png, rotate_to_depot
except ImportError:  # kalau dijalankan dengan cara lain
    from app.components.map_plot import save_route_png, rotate_to_depot


def route_table(cities, route):
    """DataFrame urutan kunjungan + jarak ke titik berikutnya (rute tertutup)."""
    r = rotate_to_depot(route)
    rows = []
    for k, idx in enumerate(r):
        nxt = r[(k + 1) % len(r)]
        a, b = cities[idx], cities[nxt]
        rows.append({
            "Urutan": k + 1,
            "Nama": a["name"],
            "X": a["x"],
            "Y": a["y"],
            "Jarak ke Berikutnya": round(math.dist((a["x"], a["y"]), (b["x"], b["y"])), 4),
        })
    first = cities[r[0]]  # baris penutup: kembali ke depot
    rows.append({"Urutan": len(r) + 1, "Nama": first["name"],
                 "X": first["x"], "Y": first["y"], "Jarak ke Berikutnya": None})
    return pd.DataFrame(rows)


def export_results(path, cities, route, total_dist, history, method_name,
                   save_png=False):
    """Simpan hasil ke `path` (.xlsx -> 3 sheet, .csv -> tabel rute saja)."""
    ext = os.path.splitext(path)[1].lower()
    table = route_table(cities, route)

    if ext == ".csv":
        table.to_csv(path, index=False)
    elif ext in (".xlsx", ".xls"):
        r = rotate_to_depot(route)
        summary = pd.DataFrame({
            "Keterangan": ["Metode", "Jumlah Kota", "Total Jarak", "Rute"],
            "Nilai": [method_name, len(cities), round(float(total_dist), 4),
                      " -> ".join(cities[i]["name"] for i in r + [r[0]])],
        })
        from openpyxl.styles import Alignment, Font, PatternFill
        head_fill = PatternFill("solid", fgColor="FF6B4A")
        head_font = Font(bold=True, color="FFFFFF")
        with pd.ExcelWriter(path, engine="openpyxl") as w:
            summary.to_excel(w, sheet_name="Ringkasan", index=False)
            table.to_excel(w, sheet_name="Rute", index=False)
            if history:
                pd.DataFrame(history).to_excel(w, sheet_name="Konvergensi", index=False)
            for ws in w.book.worksheets:  # header berwarna + lebar kolom otomatis
                for cell in ws[1]:
                    cell.fill, cell.font = head_fill, head_font
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                ws.freeze_panes = "A2"
                for col in ws.columns:
                    width = max(len(str(c.value)) if c.value is not None else 0
                                for c in col)
                    ws.column_dimensions[col[0].column_letter].width = min(width + 2, 70)
    else:
        raise ValueError(f"Format '{ext}' tidak didukung. Pakai .xlsx atau .csv")

    if save_png:
        coords = [(c["x"], c["y"]) for c in cities]
        names = [c["name"] for c in cities]
        save_route_png(os.path.splitext(path)[0] + ".png", coords, names, route,
                       title=f"Rute {method_name}",
                       subtitle=f"Total jarak: {float(total_dist):.4f}")
    return path
